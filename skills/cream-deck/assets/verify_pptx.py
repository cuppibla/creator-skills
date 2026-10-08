"""Check the editable deck against the layout it was built from, and draw a few slides back.

    python3 tools/verify_pptx.py            data check on all 56 slides
    python3 tools/verify_pptx.py 1 3 12     also re-draw those slides from the PPTX itself,
                                            under the browser screenshot, into build/verify/

The re-draw reads only what the PPTX actually contains — shape by shape, run by run — so a
builder bug shows up as a difference from the real deck above it.
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "build"
EMU_PX = 7620
_IDX: dict[tuple[str, str], int] = {}


def face(mono: bool, bold: bool, italic: bool, size: float) -> ImageFont.FreeTypeFont:
    path = "/System/Library/Fonts/Menlo.ttc" if mono else "/System/Library/Fonts/HelveticaNeue.ttc"
    want = "Bold Italic" if bold and italic else "Bold" if bold else "Italic" if italic else "Regular"
    key = (path, want)
    if key not in _IDX:
        found = 0
        for i in range(14):
            try:
                f = ImageFont.truetype(path, 20, index=i)
            except Exception:
                break
            if (f.getname()[1] or "").strip().lower() == want.lower():
                found = i
                break
        _IDX[key] = found
    return ImageFont.truetype(path, max(1, int(round(size))), index=_IDX[key])


def emu(v) -> float:
    return (v or 0) / EMU_PX


def rgb(obj) -> str | None:
    """RGBColor is a tuple subclass — "%s" formatting on it raises, so str() it."""
    try:
        return "#" + str(obj.rgb)
    except Exception:
        return None


# ── data check ──────────────────────────────────────────────────────────────
def check(prs: Presentation, layout: list[dict]) -> list[str]:
    problems: list[str] = []
    for i, (slide, lay) in enumerate(zip(prs.slides, layout), start=1):
        got = [(round(emu(s.left)), round(emu(s.top)), round(emu(s.width)), round(emu(s.height))) for s in slide.shapes]
        for el in lay["boxes"] + lay["images"] + lay["texts"]:
            want = (round(el["x"]), round(el["y"]))
            if not any(abs(want[0] - g[0]) <= 2 and abs(want[1] - g[1]) <= 2 for g in got):
                problems.append(f"slide {i}: nothing placed at {want} (source element)")
                break
        src = " ".join(" ".join(r.get("t", " ") for r in t["runs"]) for t in lay["texts"]).split()
        out = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame).split()
        miss = [w for w in src if len(w) > 3 and w not in out]
        if miss:
            problems.append(f"slide {i}: text not in the PPTX: {miss[:3]}")
        if not slide.has_notes_slide or len(slide.notes_slide.notes_text_frame.text) < 60:
            problems.append(f"slide {i}: speaker notes missing or thin")
    return problems


# ── re-draw what the PPTX contains ──────────────────────────────────────────
def render(slide) -> Image.Image:
    im = Image.new("RGB", (1600, 900), "#FAF6F0")
    d = ImageDraw.Draw(im, "RGBA")
    for sh in slide.shapes:
        x, y, w, h = emu(sh.left), emu(sh.top), emu(sh.width), emu(sh.height)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                pic = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
                cl, cr = sh.crop_left or 0, sh.crop_right or 0
                ct, cb = sh.crop_top or 0, sh.crop_bottom or 0
                if any((cl, cr, ct, cb)):
                    W0, H0 = pic.size
                    pic = pic.crop((int(W0 * cl), int(H0 * ct), int(W0 * (1 - cr)), int(H0 * (1 - cb))))
                im.paste(pic.resize((max(1, int(w)), max(1, int(h)))), (int(x), int(y)))
            except Exception as e:
                d.rectangle([x, y, x + w, y + h], outline="#C96442", width=3)
                d.text((x + 10, y + 10), f"[image: {e}]"[:70], fill="#C96442", font=face(True, False, False, 15))
            continue
        fc = rgb(sh.fill.fore_color) if (hasattr(sh, "fill") and sh.fill.type == 1) else None
        lc = rgb(sh.line.color)
        lw = max(1, int(emu(sh.line.width))) if sh.line.width else 0
        if fc or (lc and lw):
            if w < 3 or h < 3:
                d.line([x, y, x + max(w, 1), y + max(h, 1)], fill=lc or fc, width=lw or 2)
            else:
                d.rounded_rectangle([x, y, x + w, y + h], radius=16, fill=fc, outline=lc if lw else None, width=lw or 1)
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        cy = y
        for p in sh.text_frame.paragraphs:
            lh = p.line_spacing.pt / 0.6 if hasattr(p.line_spacing, "pt") else 26
            runs = [(r.text, r.font) for r in p.runs]
            widths = []
            for txt, f in runs:
                ft = face("Menlo" in (f.name or ""), bool(f.bold), bool(f.italic), (f.size.pt / 0.6) if f.size else 20)
                widths.append(ft.getlength(txt))
            total = sum(widths)
            al = str(p.alignment or "")
            ox = x + (w - total) / 2 if "CENTER" in al else x + w - total if "RIGHT" in al else x
            for (txt, f), tw in zip(runs, widths):
                ft = face("Menlo" in (f.name or ""), bool(f.bold), bool(f.italic), (f.size.pt / 0.6) if f.size else 20)
                d.text((ox, cy + lh / 2), txt, font=ft, fill=rgb(f.color) or "#2A2520", anchor="lm")
                ox += tw
            cy += lh
    return im


def main() -> None:
    layout = json.loads((BUILD / "layout.json").read_text())
    prs = Presentation(str(ROOT / "W4-The-Archive-editable.pptx"))
    problems = check(prs, layout)
    print(f"data check · {len(layout)} slides · " + ("OK — every element placed, all text present, notes on every slide"
                                                     if not problems else f"{len(problems)} problems"))
    for p in problems[:12]:
        print("   ·", p)
    want = [int(a) for a in sys.argv[1:]]
    if not want:
        return
    out = BUILD / "verify"
    out.mkdir(exist_ok=True)
    slides = list(prs.slides)
    for n in want:
        mine = render(slides[n - 1])
        real = Image.open(BUILD / "shots" / f"s{n:02d}.png").convert("RGB").resize((1600, 900))
        side = Image.new("RGB", (1600, 1816), "#111")
        side.paste(real, (0, 0))
        side.paste(mine, (0, 916))
        side.save(out / f"v{n:02d}.png")
    print("re-drawn (deck on top, PPTX below):", ", ".join(f"v{n:02d}.png" for n in want))


if __name__ == "__main__":
    main()
