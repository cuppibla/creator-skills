"""Build the week-4 PPTX decks, with the talk track as speaker notes.

    python3 tools/build_pptx.py

    W4-The-Archive-editable.pptx   native text boxes, shapes and vector (SVG) diagrams
    W4-The-Archive-flat.pptx       one rendered image per slide — nothing to knock out of place

Both are 16:9 at 1600×900 CSS pixels (13.333in × 7.5in), so one CSS pixel is exactly 7620 EMU
and 0.6 pt. Geometry for the editable deck comes from build/layout.json — the real positions
read out of the browser by tools/extract_layout.py — so it lands where the HTML deck lands.
"""
from __future__ import annotations

import base64
import json
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "build"
SRC = ROOT / "img" / "src"
ICONS = Path.home() / ".claude/skills/cream-deck/assets/icons"
if not ICONS.exists():
    ICONS = ROOT.parents[1] / "Topics/agent-memory-workshop/slide-diagrams/icons"

sys.path.insert(0, str(ROOT / "tools"))
import deckdata  # noqa: E402

W, H = 1600, 900
EMU_PX = 7620
BG = "FAF6F0"
SANS, MONO = "Helvetica Neue", "Menlo"
ALIGN = {"start": PP_ALIGN.LEFT, "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}


def px(v: float) -> Emu:
    return Emu(int(round(v * EMU_PX)))


def tp(v: float) -> Pt:
    return Pt(round(v * 0.6, 1))


# ── slide plumbing ──────────────────────────────────────────────────────────
def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    return prs


def blank(prs: Presentation, bg: str | None = BG):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if bg:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(bg)
    return slide


def notes_on(slide, text: str) -> None:
    tf = slide.notes_slide.notes_text_frame
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(12)
            r.font.name = SANS


# ── drawing ─────────────────────────────────────────────────────────────────
def _alpha(color_el, alpha: float) -> None:
    if alpha >= 0.999:
        return
    a = etree.SubElement(color_el._xFill.find(qn("a:srgbClr")) if False else color_el, qn("a:alpha"))
    a.set("val", str(int(alpha * 100000)))


def add_box(slide, b: dict):
    if b.get("rule"):
        x1, y1 = px(b["x"]), px(b["y"])
        x2, y2 = px(b["x"] + b["w"]), px(b["y"] + b["h"])
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        ln.line.color.rgb = RGBColor.from_string(b["stroke"])
        ln.line.width = px(max(b["sw"], 1))
        return ln
    r = b.get("r") or 0
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,
        px(b["x"]), px(b["y"]), px(b["w"]), px(b["h"]))
    shape.shadow.inherit = False
    if r:
        adj = min(0.5, r / max(1.0, min(b["w"], b["h"])))
        shape.adjustments[0] = adj
    if b.get("fill"):
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(b["fill"])
        if b.get("alpha", 1) < 0.999:
            srgb = shape.fill.fore_color._xFill.find(qn("a:srgbClr"))
            etree.SubElement(srgb, qn("a:alpha")).set("val", str(int(b["alpha"] * 100000)))
    else:
        shape.fill.background()
    if b.get("stroke") and b.get("sw"):
        shape.line.color.rgb = RGBColor.from_string(b["stroke"])
        shape.line.width = px(b["sw"])
        if b.get("dash"):
            ln = shape.line._get_or_add_ln()
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    else:
        shape.line.fill.background()
    shape.text_frame.word_wrap = False
    return shape


def add_text(slide, t: dict):
    tb = slide.shapes.add_textbox(px(t["x"]), px(t["y"]), px(t["w"]), px(t["h"] + 6))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras, cur = [[]], None
    for r in t["runs"]:
        if r.get("br"):
            paras.append([])
            continue
        for k, chunk in enumerate(r["t"].split("\n")):
            if k:
                paras.append([])
            if chunk != "" or len(paras[-1]) == 0:
                paras[-1].append(dict(r, t=chunk))
    para_objs = [tf.paragraphs[0]] + [tf.add_paragraph() for _ in paras[1:]]
    for p, runs in zip(para_objs, paras):
        p.alignment = ALIGN.get(t.get("align", "start"), PP_ALIGN.LEFT)
        p.line_spacing = tp(t["lh"])
        p.space_before = p.space_after = Pt(0)
        if not runs:
            r = p.add_run()
            r.text = ""
            r.font.size = tp(t["runs"][0].get("sz", 20) if t["runs"] else 20)
            continue
        for spec in runs:
            run = p.add_run()
            run.text = spec["t"]
            f = run.font
            f.name = MONO if spec.get("m") else SANS
            f.size = tp(spec["sz"])
            f.bold = spec["w"] >= 600
            f.italic = bool(spec.get("i"))
            f.color.rgb = RGBColor.from_string(spec["c"])
            rPr = run._r.get_or_add_rPr()
            if spec.get("st"):
                rPr.set("strike", "sngStrike")
            if spec.get("ls"):
                rPr.set("spc", str(int(spec["ls"] * 0.6 * 100)))
    return tb


def _round_corners(pic, r: float, w: float, h: float) -> None:
    if not r:
        return
    spPr = pic._element.spPr
    for old in spPr.findall(qn("a:prstGeom")):
        spPr.remove(old)
    geom = etree.SubElement(spPr, qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    avLst = etree.SubElement(geom, qn("a:avLst"))
    gd = etree.SubElement(avLst, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {int(min(0.5, r / max(1.0, min(w, h))) * 100000)}")
    spPr.remove(geom)
    spPr.insert(list(spPr).index(spPr.find(qn("a:xfrm"))) + 1, geom)


def add_image(slide, im: dict, base: Path):
    path = base / im["src"]
    if not path.exists():
        print("   ! missing image", im["src"])
        return None
    x, y, w, h = im["x"], im["y"], im["w"], im["h"]
    with Image.open(path) as probe:          # the file is the truth about its own size
        nw, nh = probe.size
    if im.get("fit") == "contain" and nw and nh:
        k = min(w / nw, h / nh)
        cw, ch = nw * k, nh * k
        x, y, w, h = x + (w - cw) / 2, y + (h - ch) / 2, cw, ch
    pic = slide.shapes.add_picture(str(path), px(x), px(y), px(w), px(h))
    if im.get("fit") == "cover" and nw and nh:
        k = max(w / nw, h / nh)
        sw, sh = w / k / nw, h / k / nh
        pic.crop_left = pic.crop_right = max(0.0, (1 - sw) / 2)
        pic.crop_top = pic.crop_bottom = max(0.0, (1 - sh) / 2)
    _round_corners(pic, im.get("r", 0), w, h)
    return pic


# ── vector diagrams: the PNG, with the SVG attached so PowerPoint can edit it ─
SVG_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
_svg_parts: dict[str, Part] = {}


def inlined_svg(stem: str) -> bytes | None:
    f = SRC / f"{stem}.svg"
    if not f.exists():
        return None
    s = f.read_text()

    def sub(m):
        p = ICONS / f"icon-{m.group(1)}.png"
        if not p.exists():
            return m.group(0)
        return "data:image/png;base64," + base64.b64encode(p.read_bytes()).decode()

    return re.sub(r"@@ICON:([a-z]+)@@", sub, s).encode()


def attach_svg(prs: Presentation, slide, pic, stem: str) -> bool:
    blob = inlined_svg(stem)
    if blob is None:
        return False
    part = _svg_parts.get(stem)
    if part is None:
        n = len(_svg_parts) + 1
        part = Part(PackURI(f"/ppt/media/vector{n}.svg"), "image/svg+xml", prs.part.package, blob)
        _svg_parts[stem] = part
    rId = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.blipFill.find(qn("a:blip"))
    ext_lst = blip.find(qn("a:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(blip, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"))
    ext.set("uri", SVG_EXT)
    ext.append(etree.fromstring(
        f'<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
        f'{nsdecls("r")} r:embed="{rId}"/>'))
    return True


# ── chrome ──────────────────────────────────────────────────────────────────
def add_chrome(slide, s: dict, n: int, total: int) -> None:
    tag = f"{s['act']}  ·  {s['tag']}" if s["act"] else s["tag"]
    add_text(slide, {"x": 28, "y": 862, "w": 700, "h": 20, "align": "start", "lh": 20,
                     "runs": [{"t": tag, "sz": 14, "w": 700, "c": "C9BCA9", "m": True, "ls": 2}]})
    add_text(slide, {"x": 900, "y": 862, "w": 672, "h": 20, "align": "right", "lh": 20,
                     "runs": [{"t": f"{n} / {total}", "sz": 15, "w": 700, "c": "C9BCA9", "m": True, "ls": 1}]})
    mem = s.get("mem")
    if mem:
        w = 26 + len(mem) * 9.4
        add_box(slide, {"x": (W - w) / 2, "y": 16, "w": w, "h": 29, "fill": "FFFFFF", "alpha": 0.88,
                        "stroke": "E5DCC9", "sw": 1.5, "r": 14})
        add_text(slide, {"x": (W - w) / 2, "y": 22, "w": w, "h": 18, "align": "center", "lh": 18,
                         "runs": [{"t": mem, "sz": 14, "w": 700, "c": "B9AE9B", "m": True, "ls": 2}]})


# ── the two decks ───────────────────────────────────────────────────────────
def build_editable(slides, notes, layout) -> Path:
    prs = new_deck()
    for i, s in enumerate(slides, start=1):
        lay = layout[i - 1]
        sl = blank(prs)
        items = ([("box", b) for b in lay["boxes"]] + [("img", m) for m in lay["images"]]
                 + [("txt", t) for t in lay["texts"]])
        items.sort(key=lambda kv: kv[1].get("z", 0))
        for kind, el in items:
            if kind == "box":
                add_box(sl, el)
            elif kind == "img":
                pic = add_image(sl, el, ROOT)
                stem = Path(el["src"]).stem
                if pic is not None and el["src"].startswith("img/c") and attach_svg(prs, sl, pic, stem):
                    pass
            else:
                add_text(sl, el)
        add_chrome(sl, s, i, len(slides))
        notes_on(sl, deckdata.note_text(i, notes[i]))
    out = ROOT / "W4-The-Archive-editable.pptx"
    prs.save(str(out))
    return out


def build_flat(slides, notes) -> Path:
    prs = new_deck()
    for i, s in enumerate(slides, start=1):
        shot = BUILD / "shots" / f"s{i:02d}.png"
        sl = blank(prs)
        if shot.exists():
            sl.shapes.add_picture(str(shot), 0, 0, px(W), px(H))
        else:
            print("   ! missing shot", shot.name)
        notes_on(sl, deckdata.note_text(i, notes[i]))
    out = ROOT / "W4-The-Archive-flat.pptx"
    prs.save(str(out))
    return out


def main() -> None:
    slides, notes = deckdata.slides(), deckdata.notes()
    layout = json.loads((BUILD / "layout.json").read_text())
    assert len(slides) == len(layout) == len(notes), (len(slides), len(layout), len(notes))
    a = build_editable(slides, notes, layout)
    print(f"editable · {a.relative_to(ROOT)} · {a.stat().st_size // 1024} KB · svg diagrams: {len(_svg_parts)}")
    b = build_flat(slides, notes)
    print(f"flat     · {b.relative_to(ROOT)} · {b.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
